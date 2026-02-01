import fitz, os, re, aiohttp, yt_dlp
from docx import Document

class Toolkit:
    @staticmethod
    def format_ts(ms):
        td = int(ms / 1000)
        m, s = divmod(td, 60)
        h, m = divmod(m, 60)
        return f"{h:02d}:{m:02d}:{s:02d}" if h > 0 else f"{m:02d}:{s:02d}"

    @staticmethod
    async def process_video(url):
        opts = {'skip_download': True, 'writesubtitles': True, 'writeautomaticsub': True, 'subtitleslangs': ['ru', 'en'], 'quiet': True}
        with yt_dlp.YoutubeDL(opts) as ydl:
            info = ydl.extract_info(url, download=False)
            subs = info.get('subtitles') or info.get('automatic_captions')
            if subs:
                lang = 'ru' if 'ru' in subs else 'en'
                # Ищем json3 формат для точных таймкодов
                try:
                    url_json = next(s['url'] for s in subs[lang] if 'json3' in s['url'] or 'json' in s['url'])
                    async with aiohttp.ClientSession() as s:
                        async with s.get(url_json) as r:
                            data = await r.json()
                            res = []
                            last_t = -60000
                            for ev in data.get('events', []):
                                start = ev.get('tStartMs', 0)
                                if start - last_t > 60000: # Таймкод каждую минуту
                                    res.append(f"\n📍 [{Toolkit.format_ts(start)}]")
                                    last_t = start
                                for seg in ev.get('segs', []):
                                    res.append(seg.get('utf8', '').strip())
                            return " ".join(res)
                except: pass
            return "NEED_WHISPER"

    @staticmethod
    async def parse_file(path):
        ext = path.split('.')[-1].lower()
        if ext == 'pdf':
            with fitz.open(path) as doc: return "".join([p.get_text() for p in doc])
        if ext in ['docx', 'doc']:
            return "\n".join([p.text for p in Document(path).paragraphs])
        return ""

    @staticmethod
    async def get_cloud_link(url):
        if "yadi.sk" in url or "disk.yandex" in url:
            api = f"https://cloud-api.yandex.net/v1/disk/public/resources/download?public_key={url}"
            async with aiohttp.ClientSession() as s:
                async with s.get(api) as r:
                    if r.status == 200: return (await r.json()).get("href")
        return None

    @staticmethod
    def export_file(text, fmt, uid):
        path = f"data/res_{uid}.{fmt}"
        if fmt == 'docx':
            doc = Document(); doc.add_paragraph(text); doc.save(path)
        else:
            with open(path, 'w', encoding='utf-8') as f: f.write(text)
        return path